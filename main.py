import json
import os
import sys
import uuid
from threading import Lock

import anthropic
import requests
from docx import Document
from dotenv import load_dotenv
from flask import Flask, request, jsonify, render_template_string
from waitress import serve

from Assistant_egit import assistant_instructions

load_dotenv()

MODEL = "claude-haiku-5-5"  # hesabında erişebildiğin güncel bir model adı yaz
PORT = 8000

AIRTABLE_API_KEY = os.getenv("AIRTABLE_API_KEY")
AIRTABLE_BASE_ID = os.getenv("AIRTABLE_BASE_ID")
AIRTABLE_TABLE = os.getenv("AIRTABLE_TABLE", "Leads")

app = Flask(__name__)

# ANTHROPIC_API_KEY ortam değişkenini otomatik okur
client = anthropic.Anthropic()

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
LEADS_PATH = os.path.join(BASE_DIR, "leads.json")
KNOWLEDGE_PATH = os.path.join(BASE_DIR, "knowledge.docx")

if AIRTABLE_API_KEY and AIRTABLE_BASE_ID:
    print(f"Airtable bağlantısı ayarlı (tablo: {AIRTABLE_TABLE}).")
else:
    print("UYARI: Airtable ayarları .env içinde eksik, kayıtlar sadece leads.json'a yazılacak.")


def save_lead_locally(lead):
    """Yedek olarak leads.json dosyasına yazar."""
    leads = []
    if os.path.exists(LEADS_PATH):
        with open(LEADS_PATH, encoding="utf-8") as f:
            leads = json.load(f)
    leads.append(lead)
    with open(LEADS_PATH, "w", encoding="utf-8") as f:
        json.dump(leads, f, ensure_ascii=False, indent=2)


def create_lead(name, company_name, phone, email):
    """Müşteri adayını Airtable'a kaydeder; olmazsa leads.json'a yazar."""
    lead = {
        "Name": name,
        "Company Name": company_name,
        "Phone": phone,
        "Email": email,
    }

    if AIRTABLE_API_KEY and AIRTABLE_BASE_ID:
        r = None
        try:
            r = requests.post(
                f"https://api.airtable.com/v0/{AIRTABLE_BASE_ID}/{AIRTABLE_TABLE}",
                headers={"Authorization": f"Bearer {AIRTABLE_API_KEY}"},
                json={"fields": lead, "typecast": True},
                timeout=15,
            )
            r.raise_for_status()
            print("Lead Airtable'a kaydedildi.")
            return {"status": "success"}
        except Exception as e:
            detail = r.text if r is not None else ""
            print(f"Airtable hatası: {e} {detail}")
            save_lead_locally(lead)
            print("Lead yedek olarak leads.json dosyasına kaydedildi.")
            return {"status": "success"}

    save_lead_locally(lead)
    print("Lead leads.json dosyasına kaydedildi.")
    return {"status": "success"}


def load_knowledge(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]

    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))

    return "\n".join(parts)


if not os.path.exists(KNOWLEDGE_PATH):
    print(f"HATA: Dosya bulunamadı: {KNOWLEDGE_PATH}")
    print("Dosyanın adı tam olarak 'knowledge.docx' olmalı ve main.py ile aynı klasörde durmalı.")
    sys.exit(1)

KNOWLEDGE = load_knowledge(KNOWLEDGE_PATH)
print(f"Bilgi tabanı yüklendi ({len(KNOWLEDGE)} karakter).")

SYSTEM_PROMPT = f"""{assistant_instructions}

Aşağıdaki bilgi tabanını kullanarak HanifeAkademi ile ilgili soruları cevapla:

<bilgi_tabani>
{KNOWLEDGE}
</bilgi_tabani>"""

TOOLS = [
    {
        "name": "create_lead",
        "description": "Müşteri adayı (lead) kaydı oluşturur. Ad, şirket adı ve e-posta zorunludur; telefon isteğe bağlıdır.",
        "input_schema": {
            "type": "object",
            "properties": {
                "name": {"type": "string", "description": "Kişinin adı"},
                "company_name": {"type": "string", "description": "Şirket adı"},
                "email": {"type": "string", "description": "E-posta adresi"},
                "phone": {"type": "string", "description": "Telefon numarası (verilmediyse boş metin)"},
            },
            "required": ["name", "company_name", "email"],
        },
    }
]

# thread_id -> mesaj listesi (bellekte tutulur)
conversations = {}
lock = Lock()

PAGE = """
<!doctype html>
<html lang="tr">
<head>
  <meta charset="utf-8">
  <title>Sohbet</title>
  <style>
    body { font-family: sans-serif; max-width: 600px; margin: 40px auto; }
    #log { border: 1px solid #ccc; padding: 10px; height: 350px; overflow-y: auto; margin-bottom: 10px; }
    .user { color: #0a58ca; margin: 6px 0; }
    .bot { color: #222; margin: 6px 0; }
    input { width: 80%; padding: 8px; }
    button { padding: 8px 14px; }
  </style>
</head>
<body>
  <h2>HanifeAkademi Asistanı</h2>
  <div id="log"></div>
  <input id="msg" placeholder="Mesajını yaz..." autofocus>
  <button onclick="send()">Gönder</button>

  <script>
    let threadId = null;
    const log = document.getElementById("log");
    const input = document.getElementById("msg");

    async function start() {
      const r = await fetch("/start");
      threadId = (await r.json()).thread_id;
    }

    function add(cls, text) {
      const d = document.createElement("div");
      d.className = cls;
      d.textContent = text;
      log.appendChild(d);
      log.scrollTop = log.scrollHeight;
    }

    async function send() {
      const text = input.value.trim();
      if (!text) return;
      input.value = "";
      add("user", "Sen: " + text);
      const r = await fetch("/chat", {
        method: "POST",
        headers: {"Content-Type": "application/json"},
        body: JSON.stringify({thread_id: threadId, message: text})
      });
      const data = await r.json();
      add("bot", "Bot: " + (data.response || data.error));
    }

    input.addEventListener("keydown", e => { if (e.key === "Enter") send(); });
    start();
  </script>
</body>
</html>
"""


@app.route('/')
def home():
    return render_template_string(PAGE)


@app.route('/start', methods=['GET'])
def start_conversation():
    thread_id = str(uuid.uuid4())
    with lock:
        conversations[thread_id] = []
    print(f"New thread created with ID: {thread_id}")
    return jsonify({"thread_id": thread_id})


@app.route('/chat', methods=['POST'])
def chat():
    data = request.json or {}
    thread_id = data.get('thread_id')
    user_input = data.get('message', '')

    if not thread_id:
        return jsonify({"error": "Missing thread_id"}), 400

    if thread_id not in conversations:
        return jsonify({"error": "Unknown thread_id"}), 404

    print(f"Received message: {user_input} for thread ID: {thread_id}")

    messages = conversations[thread_id]
    messages.append({"role": "user", "content": user_input})

    try:
        for _ in range(5):  # sonsuz döngüye karşı güvenlik sınırı
            response = client.messages.create(
                model=MODEL,
                max_tokens=1000,
                system=SYSTEM_PROMPT,
                tools=TOOLS,
                messages=messages,
            )

            messages.append({"role": "assistant", "content": response.content})

            if response.stop_reason != "tool_use":
                break

            tool_results = []
            for block in response.content:
                if block.type == "tool_use" and block.name == "create_lead":
                    args = block.input
                    output = create_lead(
                        args.get('name', ''),
                        args.get('company_name', ''),
                        args.get('phone', ''),
                        args.get('email', ''),
                    )
                    tool_results.append({
                        "type": "tool_result",
                        "tool_use_id": block.id,
                        "content": json.dumps(output),
                    })

            messages.append({"role": "user", "content": tool_results})

    except anthropic.APIError as e:
        print(f"Anthropic API error: {e}")
        return jsonify({"error": "Model request failed"}), 500

    reply = "".join(b.text for b in response.content if b.type == "text")

    print(f"Assistant response: {reply}")
    return jsonify({"response": reply})


if __name__ == '__main__':
    print(f"Sunucu başlıyor: http://localhost:{PORT}")
    serve(app, host='0.0.0.0', port=PORT)